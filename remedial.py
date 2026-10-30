"""ระบบสอบแก้ตัว (0 / ร / มส / มผ) — ครูส่งผลออนไลน์ นายทะเบียนตรวจรับ (public/remedial.html)

รายการ (rm_cases) = นักเรียน 1 คน × รายวิชา 1 วิชา × ภาคเรียนที่ติด  เข้าระบบได้ 2 ทาง
  ครูผู้สอนเพิ่มเอง   เลือกวิชาที่สอนในตารางสอนภาคเรียนนั้น → ติ๊กนักเรียนในชั้น + ผลเดิม
  นายทะเบียนนำเข้า   Excel / วางข้อความจาก SGS (หัวคอลัมน์ เลขประจำตัว · รหัสวิชา · ผลการเรียน [· ชื่อวิชา · ครูผู้สอน])
สถานะ  pending ค้าง → submitted ครูส่งผลแล้ว → received นายทะเบียนตรวจรับ
       นายทะเบียนตีกลับได้ (กลับเป็น pending + return_note ให้ครูเห็น)
ผลใหม่ตามระเบียบการวัดผล: 0 / มส ผ่าน = 1 (ไม่เกิน 1) · ร ได้ตามคะแนนจริง 1–4 · มผ → ผ
บัญชีครูส่วนใหญ่เป็นบัญชีรวม → ครูเลือกชื่อครูผู้สอนตอนส่ง (เก็บทั้ง teacher_id และชื่อ)
เก็บเลขประจำตัว/ชื่อ/ชั้นของนักเรียนไว้ในรายการด้วย — ลบนักเรียนออกจากทะเบียนแล้วประวัติไม่หาย
สิทธิ์: ครูทุกคนเพิ่ม/ส่งผลได้ · ตรวจรับ ตีกลับ นำเข้า ดาวน์โหลด Excel ลบรายการที่ส่งแล้ว = แอดมิน (นายทะเบียน)
"""
import csv
import datetime
import io
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from timetable import bare_name, jl, term_dates

SCHEMA = """
CREATE TABLE IF NOT EXISTS rm_cases (
    id            INTEGER PRIMARY KEY AUTOINCREMENT,
    term          TEXT NOT NULL,
    student_id    INTEGER REFERENCES students(id) ON DELETE SET NULL,
    student_code  TEXT NOT NULL,
    student_name  TEXT NOT NULL DEFAULT '',
    class_label   TEXT DEFAULT '',
    code          TEXT NOT NULL,
    title         TEXT DEFAULT '',
    teacher_id    INTEGER REFERENCES tt_teachers(id) ON DELETE SET NULL,
    teacher_name  TEXT DEFAULT '',
    orig          TEXT NOT NULL,
    status        TEXT NOT NULL DEFAULT 'pending',
    result        TEXT DEFAULT '',
    exam_date     TEXT,
    attempt       INTEGER,
    score         TEXT DEFAULT '',
    note          TEXT DEFAULT '',
    return_note   TEXT DEFAULT '',
    source        TEXT DEFAULT 'teacher',
    created_by    INTEGER,
    created_at    TEXT DEFAULT (datetime('now','localtime')),
    submitted_by  INTEGER,
    submitted_at  TEXT,
    received_by   TEXT DEFAULT '',
    received_at   TEXT,
    UNIQUE(term, student_code, code)
);
CREATE INDEX IF NOT EXISTS idx_rm_term ON rm_cases(term, status);
"""

ORIGS = ('0', 'ร', 'มส', 'มผ')
R_GRADES = ('1', '1.5', '2', '2.5', '3', '3.5', '4')
STATUS_TH = {'pending': 'ค้าง', 'submitted': 'ส่งผลแล้ว', 'received': 'ตรวจรับแล้ว'}
TERM_RE = re.compile(r'[12]/25\d\d')
DATE_RE = re.compile(r'\d{4}-\d{2}-\d{2}')


def norm_code(c):
    """เลขประจำตัวนักเรียน "03880" / 3880 / "3880.0" → "3880" """
    s = str(c if c is not None else '').strip()
    s = re.sub(r'\.0+$', '', s)
    return (s.lstrip('0') or '0') if s.isdigit() else s


def norm_orig(v):
    s = re.sub(r'[\s.]', '', str(v if v is not None else ''))
    s = {'0.0': '0', '00': '0'}.get(s, s)
    return s if s in ORIGS else None


def norm_subject(c):
    return re.sub(r'\s+', '', str(c or '')).upper()[:12]


def allowed_results(orig):
    return {'0': ('1',), 'มส': ('1',), 'ร': R_GRADES, 'มผ': ('ผ',)}.get(orig, ())


def class_label(level, room):
    return f'ม.{level}/{room}' if level else ''


def init(app, get_db, login_required, current_user, get_settings):
    from flask import jsonify, request, send_file

    with get_db() as con:
        con.executescript(SCHEMA)

    def is_admin(u):
        return bool(u) and u['role'] == 'admin'

    def default_term(con, settings):
        """ภาคเรียนที่จบไปล่าสุด (สอบแก้ตัวทำช่วงต้นภาคถัดไป) ไม่มีก็ภาคเรียนล่าสุดในตารางสอน"""
        today = datetime.date.today().isoformat()
        names = [r['name'] for r in con.execute('SELECT name FROM tt_terms')]
        ended = [(term_dates(n, settings)[1], n) for n in names if term_dates(n, settings)[1] and term_dates(n, settings)[1] < today]
        if ended:
            return max(ended)[1]
        return max(names, key=lambda n: (n[2:], n[0])) if names else ''

    def teacher_ref(con, b):
        """{teacher_id | teacher_name} → (id|None, ชื่อ)"""
        tid = b.get('teacher_id')
        if tid not in (None, ''):
            try:
                r = con.execute('SELECT id, name FROM tt_teachers WHERE id=?', (int(tid),)).fetchone()
            except (TypeError, ValueError):
                r = None
            if r:
                return r['id'], r['name']
        name = re.sub(r'\s+', ' ', str(b.get('teacher_name') or '')).strip()[:80]
        if name:
            by_bare = {bare_name(t['name']): t for t in con.execute('SELECT id, name FROM tt_teachers')}
            t = by_bare.get(bare_name(name))
            if t:
                return t['id'], t['name']
        return None, name

    def case_dict(r):
        d = dict(r)
        d['status_th'] = STATUS_TH.get(d['status'], d['status'])
        d['allowed'] = list(allowed_results(d['orig']))
        return d

    def subject_title(con, code):
        r = con.execute('SELECT name FROM tt_subjects WHERE code=?', (code,)).fetchone()
        if r and r['name']:
            return r['name']
        r = con.execute("SELECT title FROM tt_lessons WHERE code=? AND title<>'' ORDER BY id DESC LIMIT 1", (code,)).fetchone()
        return r['title'] if r else ''

    def lesson_teacher(con, term, code, level, room):
        """ครูผู้สอนวิชานั้นของชั้นนั้น จากตารางสอนภาคเรียนนั้น (นำเข้าจาก SGS ที่ไม่มีคอลัมน์ครู)"""
        cls = f'{level}/{room}'
        for r in con.execute("""SELECT l.teacher_ids, l.classes FROM tt_lessons l JOIN tt_terms t ON t.id=l.term_id
                                WHERE t.name=? AND REPLACE(l.code,' ','')=?""", (term, code)):
            ids = jl(r['teacher_ids'], [])
            if ids and (cls in jl(r['classes'], []) or not level):
                t = con.execute('SELECT id, name FROM tt_teachers WHERE id=?', (ids[0],)).fetchone()
                if t:
                    return t['id'], t['name']
        return None, ''

    def student_by_code(con):
        return {norm_code(s['student_code']): s for s in con.execute(
            'SELECT id, student_code, name, class_level, room, number FROM students')}

    def bad(msg, code=400):
        return jsonify(success=False, message=msg), code

    # ── อ่าน ──────────────────────────────────────────────
    @app.get('/api/rm/meta')
    @login_required
    def rm_meta():
        u, s = current_user(), get_settings()
        with get_db() as con:
            terms = [r['name'] for r in con.execute('SELECT name FROM tt_terms')]
            known = {r['term'] for r in con.execute('SELECT DISTINCT term FROM rm_cases')}
            terms = sorted(set(terms) | known, key=lambda n: (n[2:], n[0]), reverse=True)
            teachers = [dict(r) for r in con.execute(
                'SELECT id, name FROM tt_teachers WHERE COALESCE(active,1)=1 ORDER BY sort_order, id')]
            me = con.execute('SELECT id FROM tt_teachers WHERE user_id=?', (u['id'],)).fetchone()
            waiting = con.execute("SELECT COUNT(*) FROM rm_cases WHERE status='submitted'").fetchone()[0]
            return jsonify(terms=terms, default_term=default_term(con, s), teachers=teachers,
                           my_teacher_id=me['id'] if me else None, is_admin=is_admin(u), waiting=waiting,
                           origs=list(ORIGS), r_grades=list(R_GRADES))

    @app.get('/api/rm/cases')
    @login_required
    def rm_cases():
        a = request.args
        term = a.get('term', '')
        where, params = ['term=?'], [term]
        if a.get('teacher_id'):
            where.append('teacher_id=?'); params.append(a.get('teacher_id'))
        if a.get('status') in STATUS_TH:
            where.append('status=?'); params.append(a.get('status'))
        with get_db() as con:
            rows = con.execute(f"""SELECT c.*, s.number FROM rm_cases c LEFT JOIN students s ON s.id=c.student_id
                                   WHERE {' AND '.join('c.' + w for w in where)}
                                   ORDER BY c.code, c.class_label, s.number, c.student_code""", params).fetchall()
            counts = {k: 0 for k in STATUS_TH}
            for r in con.execute('SELECT status, COUNT(*) n FROM rm_cases WHERE term=? GROUP BY status', (term,)):
                counts[r['status']] = r['n']
        return jsonify(cases=[case_dict(r) for r in rows], counts=counts)

    @app.get('/api/rm/lessons')
    @login_required
    def rm_lessons():
        """วิชาที่ครูคนนี้สอนในภาคเรียนนั้น (สำหรับเพิ่มนักเรียนที่ติด)"""
        term, tid = request.args.get('term', ''), request.args.get('teacher_id', '')
        out = []
        with get_db() as con:
            for r in con.execute("""SELECT l.id, l.code, l.title, l.classes, l.teacher_ids FROM tt_lessons l
                                    JOIN tt_terms t ON t.id=l.term_id
                                    WHERE t.name=? AND l.kind='subject' AND l.code<>'' ORDER BY l.code""", (term,)):
                if tid and int(tid) not in jl(r['teacher_ids'], []):
                    continue
                out.append(dict(id=r['id'], code=r['code'], title=r['title'] or subject_title(con, r['code']),
                                classes=jl(r['classes'], [])))
        return jsonify(lessons=out)

    @app.get('/api/rm/students')
    @login_required
    def rm_students():
        """นักเรียนในชั้นที่เลือก ?classes=1/1,2/1 (ไม่ส่งเลขบัตร/วันเกิด)"""
        out = []
        with get_db() as con:
            for c in request.args.get('classes', '').split(','):
                m = re.fullmatch(r'([1-6])/(\d+)', c.strip())
                if not m:
                    continue
                out += [dict(r, class_label=class_label(r['class_level'], r['room'])) for r in con.execute(
                    """SELECT id, number, student_code, name, class_level, room FROM students
                       WHERE class_level=? AND room=? ORDER BY number, student_code""", (int(m.group(1)), int(m.group(2))))]
        return jsonify(students=out)

    # ── ครู: เพิ่ม / แก้ / ลบ / ส่งผล ─────────────────────
    @app.post('/api/rm/cases')
    @login_required
    def rm_add():
        """{term, code, title, teacher_id|teacher_name, items:[{student_id, orig}]} — คนที่มีอยู่แล้วข้าม"""
        u, b = current_user(), request.get_json() or {}
        term, code = str(b.get('term') or ''), norm_subject(b.get('code'))
        if not TERM_RE.fullmatch(term):
            return bad('เลือกภาคเรียน')
        if not code:
            return bad('ระบุรหัสวิชา')
        items = b.get('items') or []
        if not items:
            return bad('ยังไม่ได้เลือกนักเรียน')
        added, skipped = 0, 0
        with get_db() as con:
            tid, tname = teacher_ref(con, b)
            if not tname:
                return bad('เลือกชื่อครูผู้สอน')
            title = re.sub(r'\s+', ' ', str(b.get('title') or '')).strip()[:120] or subject_title(con, code)
            for it in items:
                orig = norm_orig(it.get('orig'))
                if not orig:
                    return bad('ผลการเรียนเดิมต้องเป็น 0 ร มส หรือ มผ')
                st = con.execute('SELECT id, student_code, name, class_level, room FROM students WHERE id=?',
                                 (it.get('student_id'),)).fetchone()
                if not st:
                    return bad('ไม่พบนักเรียน')
                cur = con.execute('INSERT OR IGNORE INTO rm_cases (term, student_id, student_code, student_name, class_label, code, title,'
                                  ' teacher_id, teacher_name, orig, source, created_by) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
                                  (term, st['id'], norm_code(st['student_code']), st['name'], class_label(st['class_level'], st['room']),
                                   code, title, tid, tname, orig, 'teacher', u['id']))
                added += cur.rowcount
                skipped += 1 - cur.rowcount
        return jsonify(success=True, added=added, skipped=skipped)

    def load_case(con, cid):
        return con.execute('SELECT * FROM rm_cases WHERE id=?', (cid,)).fetchone()

    @app.put('/api/rm/cases/<int:cid>')
    @login_required
    def rm_edit(cid):
        """แก้ผลเดิม/ครูผู้สอน/ชื่อวิชา (ก่อนตรวจรับ)"""
        u, b = current_user(), request.get_json() or {}
        with get_db() as con:
            c = load_case(con, cid)
            if not c:
                return bad('ไม่พบรายการ', 404)
            if c['status'] == 'received' and not is_admin(u):
                return bad('นายทะเบียนตรวจรับแล้ว แก้ไม่ได้')
            orig = norm_orig(b['orig']) if 'orig' in b else c['orig']
            if not orig:
                return bad('ผลการเรียนเดิมต้องเป็น 0 ร มส หรือ มผ')
            if c['status'] != 'pending' and orig != c['orig'] and c['result'] not in allowed_results(orig):
                return bad('ผลที่ส่งแล้วไม่ตรงกับผลเดิมใหม่ — ถอนผลก่อนแล้วค่อยแก้')
            tid, tname = (teacher_ref(con, b) if ('teacher_id' in b or 'teacher_name' in b)
                          else (c['teacher_id'], c['teacher_name']))
            title = re.sub(r'\s+', ' ', str(b.get('title', c['title']) or '')).strip()[:120]
            con.execute('UPDATE rm_cases SET orig=?, teacher_id=?, teacher_name=?, title=? WHERE id=?',
                        (orig, tid, tname or c['teacher_name'], title, cid))
            return jsonify(success=True, case=case_dict(load_case(con, cid)))

    @app.delete('/api/rm/cases/<int:cid>')
    @login_required
    def rm_delete(cid):
        u = current_user()
        with get_db() as con:
            c = load_case(con, cid)
            if not c:
                return bad('ไม่พบรายการ', 404)
            if c['status'] != 'pending' and not is_admin(u):
                return bad('ส่งผลแล้ว ลบไม่ได้ — ถอนผลก่อน หรือแจ้งนายทะเบียน')
            con.execute('DELETE FROM rm_cases WHERE id=?', (cid,))
        return jsonify(success=True)

    @app.post('/api/rm/cases/<int:cid>/submit')
    @login_required
    def rm_submit(cid):
        """{result, exam_date, attempt, score, note, teacher_id|teacher_name} → ส่งผลผ่าน"""
        u, b = current_user(), request.get_json() or {}
        with get_db() as con:
            c = load_case(con, cid)
            if not c:
                return bad('ไม่พบรายการ', 404)
            if c['status'] == 'received':
                return bad('นายทะเบียนตรวจรับแล้ว')
            ok = allowed_results(c['orig'])
            result = str(b.get('result') or (ok[0] if len(ok) == 1 else '')).strip()
            if result not in ok:
                return bad(f"ผล {c['orig']} แก้ตัวผ่านได้ " + ' '.join(ok))
            d = str(b.get('exam_date') or '')
            if not DATE_RE.fullmatch(d):
                return bad('ระบุวันที่สอบแก้ตัว')
            try:
                datetime.date.fromisoformat(d)
            except ValueError:
                return bad('วันที่ไม่ถูกต้อง')
            if d > datetime.date.today().isoformat():
                return bad('วันที่สอบเป็นวันข้างหน้า')
            attempt = None
            if c['orig'] == '0':
                try:
                    attempt = int(b.get('attempt') or 1)
                except (TypeError, ValueError):
                    attempt = 0
                if attempt not in (1, 2):
                    return bad('สอบแก้ตัวได้ไม่เกิน 2 ครั้ง')
            score = re.sub(r'\s+', ' ', str(b.get('score') or '')).strip()[:20]
            note = re.sub(r'\s+', ' ', str(b.get('note') or '')).strip()[:300]
            tid, tname = teacher_ref(con, b) if ('teacher_id' in b or 'teacher_name' in b) else (c['teacher_id'], c['teacher_name'])
            if not tname:
                return bad('เลือกชื่อครูผู้สอน')
            con.execute("""UPDATE rm_cases SET status='submitted', result=?, exam_date=?, attempt=?, score=?, note=?,
                           teacher_id=?, teacher_name=?, return_note='', submitted_by=?,
                           submitted_at=datetime('now','localtime') WHERE id=?""",
                        (result, d, attempt, score, note, tid, tname, u['id'], cid))
            return jsonify(success=True, case=case_dict(load_case(con, cid)))

    @app.post('/api/rm/cases/<int:cid>/withdraw')
    @login_required
    def rm_withdraw(cid):
        """ครูถอนผลที่ส่งแล้ว (ก่อนตรวจรับ) กลับไปเป็นค้าง"""
        with get_db() as con:
            c = load_case(con, cid)
            if not c:
                return bad('ไม่พบรายการ', 404)
            if c['status'] != 'submitted':
                return bad('ถอนได้เฉพาะรายการที่ส่งแล้วและยังไม่ตรวจรับ')
            con.execute("UPDATE rm_cases SET status='pending', result='', exam_date=NULL, attempt=NULL, score='',"
                        " submitted_by=NULL, submitted_at=NULL WHERE id=?", (cid,))
            return jsonify(success=True, case=case_dict(load_case(con, cid)))

    # ── นายทะเบียน ────────────────────────────────────────
    def admin_ids():
        u, b = current_user(), request.get_json() or {}
        if not is_admin(u):
            return None, None, bad('เฉพาะนายทะเบียน (แอดมิน)', 403)
        try:
            ids = [int(i) for i in (b.get('ids') or [])]
        except (TypeError, ValueError):
            return None, None, bad('รายการไม่ถูกต้อง')
        if not ids:
            return None, None, bad('ยังไม่ได้เลือกรายการ')
        return u, b, ids

    @app.post('/api/rm/receive')
    @login_required
    def rm_receive():
        u, b, ids = admin_ids()
        if u is None:
            return ids
        with get_db() as con:
            n = con.execute(f"""UPDATE rm_cases SET status='received', received_by=?, received_at=datetime('now','localtime')
                                WHERE status='submitted' AND id IN ({','.join('?' * len(ids))})""",
                            [u['full_name'], *ids]).rowcount
        return jsonify(success=True, count=n)

    @app.post('/api/rm/unreceive')
    @login_required
    def rm_unreceive():
        u, b, ids = admin_ids()
        if u is None:
            return ids
        with get_db() as con:
            n = con.execute(f"""UPDATE rm_cases SET status='submitted', received_by='', received_at=NULL
                                WHERE status='received' AND id IN ({','.join('?' * len(ids))})""", ids).rowcount
        return jsonify(success=True, count=n)

    @app.post('/api/rm/return')
    @login_required
    def rm_return():
        """ตีกลับให้ครูแก้ {ids, note} → ค้าง + ข้อความถึงครู (ผลเดิมที่ส่งมาล้างออก)"""
        u, b, ids = admin_ids()
        if u is None:
            return ids
        note = re.sub(r'\s+', ' ', str(b.get('note') or '')).strip()[:300]
        if not note:
            return bad('ระบุเหตุผลที่ตีกลับ')
        with get_db() as con:
            n = con.execute(f"""UPDATE rm_cases SET status='pending', return_note=?, received_by='', received_at=NULL,
                                result='', exam_date=NULL, attempt=NULL, score='', submitted_by=NULL, submitted_at=NULL
                                WHERE status IN ('submitted','received') AND id IN ({','.join('?' * len(ids))})""",
                            [note, *ids]).rowcount
        return jsonify(success=True, count=n)

    # ── นำเข้าจาก SGS (Excel / วางข้อความ) ─────────────────
    HEADS = {
        'student_code': ('เลขประจำตัวนักเรียน', 'เลขประจำตัว', 'รหัสนักเรียน', 'เลขที่ประจำตัว', 'student_code', 'studentid'),
        'code': ('รหัสวิชา', 'รหัสรายวิชา', 'subject_code', 'subject'),
        'orig': ('ผลการเรียน', 'ระดับผลการเรียน', 'เกรด', 'ผลการเรียนเดิม', 'grade'),
        'title': ('ชื่อวิชา', 'ชื่อรายวิชา', 'รายวิชา'),
        'teacher': ('ครูผู้สอน', 'ผู้สอน', 'ครู', 'teacher'),
        'term': ('ภาคเรียน', 'ภาคเรียน/ปีการศึกษา', 'term'),
    }

    def find_header(rows):
        """หาแถวหัวตาราง (อยู่ใน 15 แถวแรก) → (index, {field: col})"""
        def key(v):
            return re.sub(r'\s+', '', str(v or '')).lower()
        for i, row in enumerate(rows[:15]):
            cols = {}
            for j, v in enumerate(row):
                k = key(v)
                for f, names in HEADS.items():
                    if f not in cols and any(k == key(n) for n in names):
                        cols[f] = j
            if {'student_code', 'code', 'orig'} <= set(cols):
                return i, cols
        return None, None

    def read_rows():
        f = request.files.get('file')
        if f:
            name = (f.filename or '').lower()
            if name.endswith('.xlsx'):
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                rows = [list(r) for ws in wb.worksheets for r in ws.iter_rows(values_only=True)]
                wb.close()
                return rows
            if name.endswith('.csv'):
                return list(csv.reader(io.StringIO(f.read().decode('utf-8-sig', 'replace'))))
            raise ValueError('รองรับไฟล์ .xlsx หรือ .csv')
        text = request.form.get('text') or (request.get_json(silent=True) or {}).get('text') or ''
        return [line.split('\t') for line in text.splitlines() if line.strip()]

    @app.post('/api/rm/import')
    @login_required
    def rm_import():
        """นายทะเบียนนำเข้ารายชื่อที่ติด · form: file หรือ text, term, preview=1 (ดูก่อน ยังไม่บันทึก)"""
        u = current_user()
        if not is_admin(u):
            return bad('เฉพาะนายทะเบียน (แอดมิน)', 403)
        args = request.form if request.files or request.form else (request.get_json(silent=True) or {})
        term = str(args.get('term') or '')
        preview = str(args.get('preview') or '') in ('1', 'true')
        if not TERM_RE.fullmatch(term):
            return bad('เลือกภาคเรียน')
        try:
            rows = read_rows()
        except Exception as e:                           # noqa: BLE001 — ไฟล์เสีย/ผิดชนิด
            return bad(f'เปิดไฟล์ไม่ได้: {e}')
        hi, cols = find_header(rows)
        if hi is None:
            return bad('ไม่พบหัวตาราง ต้องมีคอลัมน์ เลขประจำตัว · รหัสวิชา · ผลการเรียน')
        items, unknown, skipped_grade = [], [], 0
        with get_db() as con:
            studs = student_by_code(con)
            seen = set()
            for row in rows[hi + 1:]:
                cell = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None   # noqa: E731
                sc, code = norm_code(cell('student_code')), norm_subject(cell('code'))
                if not sc or not code:
                    continue
                row_term = str(cell('term') or '').strip()
                if row_term and TERM_RE.fullmatch(row_term) and row_term != term:
                    continue
                orig = norm_orig(cell('orig'))
                if not orig:
                    skipped_grade += 1                   # ผลผ่านแล้ว (1–4 ผ) ไม่ต้องแก้
                    continue
                if (sc, code) in seen:
                    continue
                seen.add((sc, code))
                st = studs.get(sc)
                if not st:
                    unknown.append(sc)
                    continue
                title = re.sub(r'\s+', ' ', str(cell('title') or '')).strip()[:120] or subject_title(con, code)
                if cell('teacher'):
                    tid, tname = teacher_ref(con, {'teacher_name': cell('teacher')})
                else:
                    tid, tname = lesson_teacher(con, term, code, st['class_level'], st['room'])
                exists = con.execute('SELECT 1 FROM rm_cases WHERE term=? AND student_code=? AND code=?',
                                     (term, sc, code)).fetchone()
                items.append(dict(student_id=st['id'], student_code=sc, student_name=st['name'],
                                  class_label=class_label(st['class_level'], st['room']), code=code, title=title,
                                  teacher_id=tid, teacher_name=tname, orig=orig, exists=bool(exists)))
            added = 0
            if not preview:
                for it in items:
                    if it['exists']:
                        continue
                    added += con.execute('INSERT OR IGNORE INTO rm_cases (term, student_id, student_code, student_name, class_label,'
                                         ' code, title, teacher_id, teacher_name, orig, source, created_by)'
                                         ' VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
                                         (term, it['student_id'], it['student_code'], it['student_name'], it['class_label'],
                                          it['code'], it['title'], it['teacher_id'], it['teacher_name'], it['orig'],
                                          'sgs', u['id'])).rowcount
        return jsonify(success=True, preview=preview, items=items, added=added, unknown=sorted(set(unknown)),
                       passed_rows=skipped_grade,
                       new=sum(1 for it in items if not it['exists']), existing=sum(1 for it in items if it['exists']))

    # ── ดาวน์โหลด Excel (นายทะเบียนนำไปกรอก SGS) ──────────
    @app.get('/api/rm/export.xlsx')
    @login_required
    def rm_export():
        u = current_user()
        if not is_admin(u):
            return bad('เฉพาะนายทะเบียน (แอดมิน)', 403)
        term = request.args.get('term', '')
        with get_db() as con:
            rows = con.execute("""SELECT c.*, s.number FROM rm_cases c LEFT JOIN students s ON s.id=c.student_id
                                  WHERE c.term=? ORDER BY c.class_label, s.number, c.student_code, c.code""", (term,)).fetchall()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'สอบแก้ตัว'
        head = ['ชั้น', 'เลขที่', 'เลขประจำตัว', 'ชื่อ-สกุล', 'รหัสวิชา', 'ชื่อวิชา', 'ครูผู้สอน', 'ผลเดิม',
                'ผลใหม่', 'ครั้งที่', 'คะแนน', 'วันที่สอบ', 'สถานะ', 'ตรวจรับโดย', 'หมายเหตุ']
        ws.append([f'ผลการสอบแก้ตัว ภาคเรียนที่ {term}'])
        ws.append(head)
        for r in rows:
            d = r['exam_date']
            dt = f'{int(d[8:10])}/{int(d[5:7])}/{int(d[:4]) + 543}' if d else ''
            ws.append([r['class_label'], r['number'], r['student_code'], r['student_name'], r['code'], r['title'],
                       r['teacher_name'], r['orig'], r['result'], r['attempt'], r['score'], dt,
                       STATUS_TH.get(r['status'], r['status']), r['received_by'], r['note'] or r['return_note']])
        f, b = Font(name='TH SarabunPSK', size=16), Font(name='TH SarabunPSK', size=16, bold=True)
        thin = Side(style='thin')
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for c in row:
                c.font = b if c.row <= 2 else f
                if c.row >= 2:
                    c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                    c.alignment = Alignment(vertical='center', horizontal='left' if c.column in (4, 6, 7, 15) else 'center')
        for c in ws[2]:
            c.fill = PatternFill('solid', fgColor='DDEBF7')
        for col, w in zip('ABCDEFGHIJKLMNO', (8, 6, 11, 26, 10, 26, 24, 8, 8, 7, 8, 12, 12, 20, 30)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = 'A3'
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, as_attachment=True, download_name=f"สอบแก้ตัว_{term.replace('/', '-')}.xlsx",
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
